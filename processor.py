import os
import json
from openai import OpenAI

def extract_text(file_bytes: bytes) -> str:
    try:
        import pdfplumber
        import io
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            text = ""
            for page in pdf.pages:
                text += page.extract_text() or ""
        if text.strip():
            return text
    except Exception:
        pass
    try:
        import openpyxl
        import io
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        text = ""
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                text += " ".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
        if text.strip():
            return text
    except Exception:
        pass
    return file_bytes.decode("utf-8", errors="ignore")

def process_file(file_bytes: bytes) -> list[dict]:
    text_content = extract_text(file_bytes)
    if not text_content.strip():
        return []
    client = OpenAI(api_key=os.environ["DEEPSEEK_API_KEY"], base_url="https://api.deepseek.com")
    system_prompt = (
        "You are an AI assistant that extracts structured job profitability data from documents. "
        "Given the extracted text, identify each job/project. For each job, output a JSON array with these fields:\n"
        "- title (string): the job or project name\n"
        "- status (string): one of exactly these values:\n"
        "    'over_budget:critical'      — already over budget\n"
        "    'margin_declining:warning'  — spending faster than expected, at risk\n"
        "    'on_track:good'             — within budget, healthy margin\n"
        "    'within_budget:good'        — under budget\n"
        "    'completed:neutral'         — job is finished\n"
        "- details (object): include all numeric data found — estimate, actual_cost, revenue, "
        "margin_percent, labor_cost, materials_cost, remaining_budget, percent_complete\n"
        "- due_date (ISO date string or null)\n"
        "- summary (string): 1-2 plain English sentences explaining the job's financial situation "
        "and WHY it is over/under budget or at risk. Be specific — mention actual numbers.\n"
        "- forecast (string): 1 sentence predicting final profitability based on current spend rate. "
        "Example: 'At current spend rate, this job will finish 15% over budget.'\n"
        "Return ONLY the JSON array, no markdown, no other text."
    )
    user_prompt = f"Document text:\n{text_content}\n\nExtract the data as specified."
    try:
        response = client.chat.completions.create(
            model="deepseek-v4-flash",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt}
            ],
            temperature=0
        )
        result_text = response.choices[0].message.content
        if not result_text:
            result_text = response.choices[0].message.reasoning_content or ""
        result_text = result_text.strip().replace("```json", "").replace("```", "").strip()
        records = json.loads(result_text)
        if isinstance(records, list):
            return [{
                "title": r.get("title", ""),
                "status": r.get("status", "on_track:good"),
                "details": r.get("details", {}),
                "due_date": r.get("due_date"),
                "summary": r.get("summary", ""),
                "forecast": r.get("forecast", "")
            } for r in records]
        return []
    except Exception as e:
        print(f"Processor error: {e}")
        return []
