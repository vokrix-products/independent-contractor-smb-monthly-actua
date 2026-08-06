import os
import json
from openai import OpenAI

def extract_text(file_bytes: bytes) -> str:
    # Try PDF
    try:
        import pdfplumber
        import io
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            text = ""
            for page in pdf.pages:
                text += page.extract_text() or ""
        if text.strip():
            return text
    except Exception:
        pass
    # Try Excel
    try:
        import openpyxl
        import io
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        text = ""
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                text += " ".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
        if text.strip():
            return text
    except Exception:
        pass
    # Fallback: decode as text
    text = file_bytes.decode("utf-8", errors="ignore")
    return text

def process_file(file_bytes: bytes) -> list[dict]:
    text_content = extract_text(file_bytes)
    if not text_content.strip():
        return []
    client = OpenAI(api_key=os.environ["DEEPSEEK_API_KEY"], base_url="https://api.deepseek.com")
    system_prompt = (
        "You are an AI assistant that extracts structured job profitability data from documents. "
        "Given the extracted text, identify each job/project. For each job, output a JSON array of objects with exactly these fields:\n"
        "- title: the job/project name (string)\n"
        "- status: must be exactly \"over_budget:critical\" if actual cost exceeds budget significantly, otherwise \"within_budget:good\"\n"
        "- details: an object containing all numeric and text details about the job (budget, actual, variance, etc.)\n"
        "- due_date: ISO-8601 date string if a relevant deadline or report date is mentioned, otherwise null\n"
        "Return ONLY the JSON array, no other text."
    )
    user_prompt = f"Document text:\n{text_content}\n\nExtract the data as specified."
    try:
        response = client.chat.completions.create(
            model="deepseek-chat",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt}
            ],
            temperature=0
        )
        result_text = response.choices[0].message.content.strip()
        # Parse JSON
        records = json.loads(result_text)
        if isinstance(records, list):
            validated = []
            for record in records:
                validated.append({
                    "title": record.get("title", ""),
                    "status": record.get("status", "within_budget:good"),
                    "details": record.get("details", {}),
                    "due_date": record.get("due_date")
                })
            return validated
        else:
            return []
    except Exception as e:
        return []
